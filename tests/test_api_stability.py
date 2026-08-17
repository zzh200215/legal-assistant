import json
import unittest
from datetime import datetime
from unittest.mock import AsyncMock, patch
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.core.auth import create_access_token
from app.core.auth import hash_password
from app.core.config import get_settings
from app.core.database import Base, get_db
from app.core import database as database_module
from app.main import app
from app.models.agent import AgentRun, ToolCallLog
from app.models.document import Document, DocumentQARecord
from app.models.llm_call_log import LLMCallLog
from app.models.operation_log import OperationLog
from app.models.prompt import PromptTemplate, PromptTemplateVersion
from app.models.task import Task
from app.models.token_usage import TokenUsage
from app.models.user import User
from app.services.llm.llm_governance_service import llm_governance_service
from app.tools.base import BaseAgentTool, tool_success


class FakeTool(BaseAgentTool):
    def __init__(self, name, description, auto_context_fields=(), handler=None, parameters=None):
        self.name = name
        self.description = description
        self.auto_context_fields = auto_context_fields
        self.parameters = parameters or {"type": "object", "properties": {}, "required": []}
        self._handler = handler or (lambda **kwargs: tool_success("ok", kwargs))

    async def run(self, **kwargs):
        return self._handler(**kwargs)


class ApiStabilityTests(unittest.TestCase):
    def _ws_subprotocols(self):
        return ["json", f"bearer.{self.token}"]

    def setUp(self):
        engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            future=True,
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        self.TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
        Base.metadata.create_all(bind=engine)
        self.db = self.TestingSessionLocal()

        self.user = User(
            username="tester",
            email="tester@example.com",
            hashed_password=hash_password("secret"),
        )
        self.db.add(self.user)
        self.db.commit()
        self.db.refresh(self.user)

        self.document = Document(
            user_id=self.user.id,
            title="spec.md",
            file_path="uploads/spec.md",
            file_type="md",
            status="indexed",
        )
        self.db.add(self.document)
        self.db.commit()
        self.db.refresh(self.document)
        self.settings = get_settings()
        self._governance_setting_keys = [
            "LLM_RATE_LIMIT_WINDOW_SECONDS",
            "LLM_RATE_LIMIT_MAX_REQUESTS",
            "LLM_DAILY_REQUEST_LIMIT",
            "LLM_DAILY_TOKEN_LIMIT",
            "LLM_ESTIMATED_CHARS_PER_TOKEN",
            "LLM_ESTIMATED_COMPLETION_TOKENS",
        ]
        self._governance_settings_backup = {
            key: getattr(self.settings, key)
            for key in self._governance_setting_keys
        }
        llm_governance_service.reset_local_state()
        self.sessionlocal_patchers = [
            patch("app.services.llm.llm_governance_service.SessionLocal", self.TestingSessionLocal),
            patch("app.services.llm.llm_observability_service.SessionLocal", self.TestingSessionLocal),
            patch("app.core.database.SessionLocal", self.TestingSessionLocal),
        ]
        for patcher in self.sessionlocal_patchers:
            patcher.start()

        def override_get_db():
            db = self.TestingSessionLocal()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = override_get_db
        self.client = TestClient(app)
        self.token = create_access_token({"sub": self.user.id})
        self.headers = {"Authorization": f"Bearer {self.token}"}

    def tearDown(self):
        app.dependency_overrides.clear()
        for patcher in reversed(getattr(self, "sessionlocal_patchers", [])):
            patcher.stop()
        for key, value in self._governance_settings_backup.items():
            setattr(self.settings, key, value)
        llm_governance_service.reset_local_state()
        self.db.close()

    def _set_governance_settings(self, **overrides):
        for key, value in overrides.items():
            setattr(self.settings, key, value)
        llm_governance_service.reset_local_state()

    def test_chat_request_validates_message_role(self):
        response = self.client.post(
            "/api/chat/",
            headers=self.headers,
            json={"messages": [{"role": "invalid", "content": "hello"}]},
        )

        self.assertEqual(response.status_code, 422)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "VALIDATION_ERROR")

    def test_chat_request_returns_stable_rate_limit_code(self):
        self._set_governance_settings(
            LLM_RATE_LIMIT_WINDOW_SECONDS=60,
            LLM_RATE_LIMIT_MAX_REQUESTS=1,
            LLM_DAILY_REQUEST_LIMIT=0,
            LLM_DAILY_TOKEN_LIMIT=0,
        )

        fake_response = {
            "choices": [{"message": {"content": "收到"}}],
            "usage": {"prompt_tokens": 12, "completion_tokens": 8},
        }
        with patch("app.core.llm_client.LLMClient._post_json_with_retry", new=AsyncMock(return_value=fake_response)):
            first = self.client.post(
                "/api/chat/",
                headers=self.headers,
                json={"messages": [{"role": "user", "content": "第一条消息"}]},
            )
            second = self.client.post(
                "/api/chat/",
                headers=self.headers,
                json={"messages": [{"role": "user", "content": "第二条消息"}]},
            )

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 429)
        payload = second.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "LLM_RATE_LIMIT_EXCEEDED")
        self.assertEqual(payload["message"], "请求过于频繁，请稍后再试")

    def test_chat_request_returns_stable_daily_token_budget_code(self):
        self._set_governance_settings(
            LLM_RATE_LIMIT_WINDOW_SECONDS=0,
            LLM_RATE_LIMIT_MAX_REQUESTS=0,
            LLM_DAILY_REQUEST_LIMIT=0,
            LLM_DAILY_TOKEN_LIMIT=100,
            LLM_ESTIMATED_CHARS_PER_TOKEN=1,
            LLM_ESTIMATED_COMPLETION_TOKENS=40,
        )
        self.db.add(
            TokenUsage(
                user_id=self.user.id,
                model="qwen-plus",
                action="chat",
                budget_category="text",
                prompt_tokens=30,
                completion_tokens=30,
                total_tokens=60,
            )
        )
        self.db.commit()

        response = self.client.post(
            "/api/chat/",
            headers=self.headers,
            json={"messages": [{"role": "user", "content": "这是一条会触发预算拦截的消息"}]},
        )

        self.assertEqual(response.status_code, 429)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "LLM_DAILY_TOKEN_BUDGET_EXCEEDED")
        self.assertEqual(payload["message"], "今日 Token 预算已用尽")

    def test_chat_request_500_does_not_leak_internal_error_detail(self):
        with patch(
            "app.api.conversation.chat_api.llm_service.chat",
            new=AsyncMock(side_effect=RuntimeError("api-key=test-secret")),
        ):
            response = self.client.post(
                "/api/chat/",
                headers=self.headers,
                json={"messages": [{"role": "user", "content": "你好"}]},
            )

        self.assertEqual(response.status_code, 500)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "CHAT_REQUEST_FAILED")
        self.assertEqual(payload["error"]["detail"], "聊天请求失败")
        self.assertNotIn("test-secret", json.dumps(payload, ensure_ascii=False))

    def test_ws_agent_rejects_invalid_max_steps(self):
        with self.client.websocket_connect("/api/ws/agent", subprotocols=self._ws_subprotocols(), timeout=5) as websocket:
            websocket.send_json({"goal": "生成计划", "max_steps": 99})
            welcome = websocket.receive_json()
            message = websocket.receive_json()

        self.assertEqual(welcome["type"], "welcome")
        self.assertEqual(message["type"], "error")
        self.assertIn("max_steps", message["message"])

    def test_ws_agent_internal_error_does_not_leak_detail(self):
        with patch(
            "app.api.conversation.ws_api.agent_service.run",
            new=AsyncMock(side_effect=RuntimeError("db_password=secret")),
        ):
            with self.client.websocket_connect("/api/ws/agent", subprotocols=self._ws_subprotocols(), timeout=5) as websocket:
                websocket.send_json({"goal": "生成计划"})
                welcome = websocket.receive_json()
                message = websocket.receive_json()

        self.assertEqual(welcome["type"], "welcome")
        self.assertEqual(message["type"], "error")
        self.assertEqual(message["message"], "请求失败")
        self.assertNotIn("secret", json.dumps(message, ensure_ascii=False))

    def test_ws_agent_can_resume_approval_run(self):
        fake_run = type(
            "FakeRun",
            (),
            {
                "id": 12,
                "user_id": self.user.id,
                "session_id": None,
                "goal": "审批恢复",
                "status": "completed",
                "result": '{"final_answer":"已恢复","artifacts":{"documents":[],"tasks":[],"emails":[],"meetings":[]}}',
                "final_answer": "已恢复",
                "last_observation": "",
                "failure_reason": None,
                "total_steps": 2,
                "error": None,
                "created_at": None,
                "completed_at": None,
            },
        )()

        with patch(
            "app.api.conversation.ws_api.agent_service.resume_after_approval",
            new=AsyncMock(return_value=fake_run),
        ), patch(
            "app.api.conversation.ws_api.agent_service.get_run_logs",
            return_value=[],
        ), patch(
            "app.api.conversation.ws_api.agent_service.serialize_run",
            return_value={"id": 12, "status": "completed", "artifacts": {}},
        ):
            with self.client.websocket_connect("/api/ws/agent", subprotocols=self._ws_subprotocols(), timeout=5) as websocket:
                websocket.send_json({"action": "resume_approval", "approval_id": 9})
                welcome = websocket.receive_json()
                message = websocket.receive_json()

        self.assertEqual(welcome["type"], "welcome")
        self.assertEqual(message["type"], "run_snapshot")
        self.assertEqual(message["run"]["id"], 12)
        self.assertEqual(message["run"]["status"], "completed")

    def test_agent_run_detail_keeps_contract_and_includes_artifacts(self):
        from app.models.agent import AgentRun

        run = AgentRun(
            user_id=self.user.id,
            goal="总结文档 1，并提取其中的风险点",
            status="completed",
            result='{"final_answer":"已完成","artifacts":{"documents":[{"document_id":1}],"tasks":[],"emails":[],"meetings":[]},"supervisor_plan":{"intent":"文档风险分析","workers":["document_agent"],"risk_level":"low"}}',
            final_answer="已完成",
            total_steps=3,
        )
        self.db.add(run)
        self.db.commit()
        self.db.refresh(run)

        response = self.client.get(f"/api/agent/runs/{run.id}", headers=self.headers)

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        data = payload["data"]
        self.assertEqual(data["final_answer"], "已完成")
        self.assertIn("artifacts", data)
        self.assertEqual(data["artifacts"]["documents"][0]["document_id"], 1)
        self.assertEqual(data["supervisor_plan"]["workers"], ["document_agent"])

    def test_ws_chat_rejects_too_long_content(self):
        with self.client.websocket_connect("/api/ws/chat", subprotocols=self._ws_subprotocols(), timeout=5) as websocket:
            websocket.send_json({"content": "a" * 8001})
            welcome = websocket.receive_json()
            message = websocket.receive_json()

        self.assertEqual(welcome["type"], "welcome")
        self.assertEqual(message["type"], "error")
        self.assertIn("消息长度不能超过", message["message"])

    def test_ws_chat_internal_error_does_not_leak_detail(self):
        async def _raise_stream(*_args, **_kwargs):
            raise RuntimeError("token=secret")
            yield

        with patch("app.api.conversation.ws_api.llm_client.chat_stream", side_effect=_raise_stream):
            with self.client.websocket_connect("/api/ws/chat", subprotocols=self._ws_subprotocols(), timeout=5) as websocket:
                websocket.send_json({"content": "你好"})
                welcome = websocket.receive_json()
                session_msg = websocket.receive_json()
                message = websocket.receive_json()

        self.assertEqual(welcome["type"], "welcome")
        self.assertEqual(session_msg["type"], "session")
        self.assertEqual(message["type"], "error")
        self.assertEqual(message["message"], "请求失败")
        self.assertNotIn("secret", json.dumps(message, ensure_ascii=False))

    def test_ws_chat_records_document_qa(self):
        fake_result = {
            "answer": "首付款应在 2026-07-01 前支付 100 万元。",
            "citations": [
                {
                    "document_id": self.document.id,
                    "chunk_id": 1,
                    "chunk_index": 0,
                    "page_number": 3,
                    "section_title": "付款条款",
                    "locator": f"doc:{self.document.id} | page:3 | section:付款条款 | chunk:0",
                    "source_text": "付款条款：甲方应在 2026-07-01 前支付首付款 100 万元。",
                }
            ],
            "confidence": 0.81,
            "can_answer": True,
            "refusal_reason": None,
            "hit_chunks": [
                {
                    "content": "付款条款：甲方应在 2026-07-01 前支付首付款 100 万元。",
                    "metadata": {"chunk_index": 0, "page_number": 3, "section_title": "付款条款"},
                }
            ],
            "latency_ms": 12,
            "agentic_rag": {"enabled": True, "retrieval_rounds": 1, "steps": []},
        }

        with patch("app.api.conversation.ws_api.agentic_rag_service.answer_async", new=AsyncMock(return_value=fake_result)):
            with self.client.websocket_connect("/api/ws/chat", subprotocols=self._ws_subprotocols(), timeout=5) as websocket:
                websocket.send_json({"content": "付款条件是什么", "document_id": self.document.id})
                welcome = websocket.receive_json()
                session_msg = websocket.receive_json()
                done = websocket.receive_json()

        self.assertEqual(welcome["type"], "welcome")
        self.assertEqual(session_msg["type"], "session")
        self.assertEqual(done["type"], "done")
        self.assertTrue(done["can_answer"])
        self.assertGreater(done["confidence"], 0.5)
        self.assertIsNone(done["refusal_reason"])
        self.assertEqual(done["agentic_rag"]["retrieval_rounds"], 1)
        self.assertEqual(done["citations"][0]["document_id"], self.document.id)
        self.assertEqual(done["citations"][0]["page_number"], 3)

        db = self.TestingSessionLocal()
        try:
            records = (
                db.query(DocumentQARecord)
                .filter_by(document_id=self.document.id, user_id=self.user.id)
                .all()
            )
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].question, "付款条件是什么")
            self.assertEqual(records[0].answer, fake_result["answer"])
            self.assertEqual(records[0].source, "ws_chat")
        finally:
            db.close()

    def test_document_and_task_list_use_paginated_structure(self):
        task = Task(
            user_id=self.user.id,
            title="跟进付款",
            status="todo",
            priority="medium",
        )
        self.db.add(task)
        self.db.commit()

        document_response = self.client.get("/api/documents/?page=1&page_size=10", headers=self.headers)
        task_response = self.client.get("/api/tasks/?page=1&page_size=10", headers=self.headers)

        self.assertEqual(document_response.status_code, 200)
        self.assertEqual(task_response.status_code, 200)

        document_payload = document_response.json()
        task_payload = task_response.json()

        self.assertTrue(document_payload["success"])
        self.assertEqual(document_payload["data"]["page"], 1)
        self.assertEqual(document_payload["data"]["page_size"], 10)
        self.assertIsInstance(document_payload["data"]["items"], list)
        self.assertGreaterEqual(document_payload["data"]["total"], 1)

        self.assertTrue(task_payload["success"])
        self.assertEqual(task_payload["data"]["page"], 1)
        self.assertEqual(task_payload["data"]["page_size"], 10)
        self.assertIsInstance(task_payload["data"]["items"], list)
        self.assertGreaterEqual(task_payload["data"]["total"], 1)

    def test_document_not_found_returns_stable_error_code(self):
        response = self.client.get("/api/documents/9999", headers=self.headers)

        self.assertEqual(response.status_code, 404)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "DOCUMENT_NOT_FOUND")
        self.assertEqual(payload["message"], "文档不存在")

    def test_document_visual_analyze_returns_expected_contract(self):
        self.document.file_type = "png"
        self.document.file_path = "uploads/spec.png"
        self.db.commit()

        with patch(
            "app.api.documents.document_api.document_service.analyze_visual",
            new=AsyncMock(
                return_value={
                    "document_id": self.document.id,
                    "title": self.document.title,
                    "file_type": "png",
                    "analysis": "图片中包含签字和公章。",
                    "image_count": 1,
                }
            ),
        ):
            response = self.client.post(
                f"/api/documents/{self.document.id}/analyze-visual",
                headers=self.headers,
                json={"prompt": "请识别图中的签字和公章"},
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["document_id"], self.document.id)
        self.assertEqual(payload["data"]["file_type"], "png")
        self.assertIn("签字和公章", payload["data"]["analysis"])

    def test_document_visual_analyze_returns_stable_error_code_for_invalid_type(self):
        with patch(
            "app.api.documents.document_api.document_service.analyze_visual",
            new=AsyncMock(side_effect=ValueError("Document visual analysis only supports image and PDF files")),
        ):
            response = self.client.post(
                f"/api/documents/{self.document.id}/analyze-visual",
                headers=self.headers,
                json={"prompt": "请识别图中的签字和公章"},
            )

        self.assertEqual(response.status_code, 400)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "DOCUMENT_VISUAL_ANALYSIS_INVALID")

    def test_document_ask_keeps_contract_when_visual_analysis_is_used(self):
        self.document.file_type = "png"
        self.document.file_path = "uploads/spec.png"
        self.db.commit()

        with patch(
            "app.api.documents.document_api.document_service.ask",
            return_value={
                "qa_record_id": 9,
                "answer": "该页包含签字和公章。",
                "citations": [],
                "confidence": 0.83,
                "can_answer": True,
                "feedback_value": None,
                "feedback_status": None,
            },
        ):
            response = self.client.post(
                f"/api/documents/{self.document.id}/ask",
                headers=self.headers,
                json={"question": "这页有没有签字"},
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["qa_record_id"], 9)
        self.assertTrue(payload["data"]["can_answer"])

    def test_invalid_credentials_return_stable_error_code(self):
        response = self.client.post(
            "/api/auth/login",
            json={"username": "tester", "password": "wrong-password"},
        )

        self.assertEqual(response.status_code, 401)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "INVALID_CREDENTIALS")

    def test_admin_only_analytics_scope_returns_stable_error_code(self):
        response = self.client.get(
            "/api/analytics/llm-calls?scope=all",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 403)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "ADMIN_REQUIRED")
        self.assertEqual(payload["message"], "需要管理员权限")

    def test_llm_call_list_hides_sensitive_fields_for_non_admin(self):
        self.db.add(
            LLMCallLog(
                user_id=self.user.id,
                module_name="document",
                action="rag_answer",
                model_name="qwen-plus",
                status="error",
                error_message="upstream timeout",
                request_excerpt="secret prompt",
                response_excerpt="secret response",
            )
        )
        self.db.commit()

        response = self.client.get("/api/analytics/llm-calls", headers=self.headers)

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        item = payload["data"]["items"][0]
        self.assertIsNone(item["error_message"])
        self.assertIsNone(item["request_excerpt"])
        self.assertIsNone(item["response_excerpt"])

    def test_health_check_does_not_leak_internal_error_message(self):
        class FakeDbSession:
            def execute(self, *_args, **_kwargs):
                raise RuntimeError("db password exposed")

            def close(self):
                return None

        class FakeRedisClient:
            def ping(self):
                raise RuntimeError("redis host exposed")

        class FakeHttpClient:
            def __init__(self, *args, **kwargs):
                pass

            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return False

            def get(self, *_args, **_kwargs):
                raise RuntimeError("llm upstream exposed")

        with patch("app.main.SessionLocal", return_value=FakeDbSession()), patch(
            "app.main.redis.from_url",
            return_value=FakeRedisClient(),
        ), patch("app.main.httpx.Client", FakeHttpClient):
            response = self.client.get("/api/health")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["status"], "degraded")
        for check in payload["data"]["checks"].values():
            self.assertEqual(set(check.keys()) & {"message", "detail", "error"}, set())

    def test_unhandled_500_does_not_leak_internal_error_detail(self):
        @app.get("/api/test-crash")
        def _test_crash():
            raise RuntimeError("db_password=super-secret")

        try:
            crash_client = TestClient(app, raise_server_exceptions=False)
            response = crash_client.get("/api/test-crash")
        finally:
            app.router.routes.pop()

        self.assertEqual(response.status_code, 500)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "INTERNAL_SERVER_ERROR")
        self.assertEqual(payload["error"]["detail"], "服务器内部错误")
        self.assertNotIn("super-secret", json.dumps(payload, ensure_ascii=False))

    def test_experiment_overview_requires_admin_permission(self):
        response = self.client.get(
            "/api/analytics/experiments/overview",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 403)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "ADMIN_REQUIRED")

    def test_document_qa_feedback_submission_returns_expected_contract(self):
        qa_record = DocumentQARecord(
            document_id=self.document.id,
            user_id=self.user.id,
            question="付款时间是什么？",
            answer="2026-07-01",
            source="document",
        )
        self.db.add(qa_record)
        self.db.commit()
        self.db.refresh(qa_record)

        response = self.client.post(
            f"/api/documents/qa-records/{qa_record.id}/feedback",
            headers=self.headers,
            json={
                "feedback_value": "negative",
                "feedback_reason": "wrong_citation",
                "feedback_note": "引用页码不对",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["id"], qa_record.id)
        self.assertEqual(payload["data"]["feedback_value"], "negative")
        self.assertEqual(payload["data"]["feedback_reason"], "wrong_citation")
        self.assertEqual(payload["data"]["feedback_status"], "open")

    def test_document_qa_feedback_requires_existing_record(self):
        response = self.client.post(
            "/api/documents/qa-records/9999/feedback",
            headers=self.headers,
            json={"feedback_value": "positive"},
        )

        self.assertEqual(response.status_code, 404)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "QA_RECORD_NOT_FOUND")

    def test_feedback_resolve_requires_admin_permission(self):
        qa_record = DocumentQARecord(
            document_id=self.document.id,
            user_id=self.user.id,
            question="付款时间是什么？",
            answer="2026-07-01",
            source="document",
            feedback_value="negative",
            feedback_status="open",
        )
        self.db.add(qa_record)
        self.db.commit()
        self.db.refresh(qa_record)

        response = self.client.post(
            f"/api/analytics/feedback/{qa_record.id}/resolve",
            headers=self.headers,
            json={"resolution_note": "已复盘"},
        )

        self.assertEqual(response.status_code, 403)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "ADMIN_REQUIRED")

    def test_admin_can_resolve_negative_feedback(self):
        admin = User(
            username="admin_feedback",
            email="admin_feedback@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        self.db.refresh(admin)
        admin_token = create_access_token({"sub": admin.id})

        qa_record = DocumentQARecord(
            document_id=self.document.id,
            user_id=self.user.id,
            question="付款时间是什么？",
            answer="2026-07-01",
            source="document",
            feedback_value="negative",
            feedback_status="open",
        )
        self.db.add(qa_record)
        self.db.commit()
        self.db.refresh(qa_record)

        response = self.client.post(
            f"/api/analytics/feedback/{qa_record.id}/resolve",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"resolution_note": "已修正 Prompt"},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["feedback_status"], "resolved")
        self.assertEqual(payload["data"]["feedback_resolution_note"], "已修正 Prompt")

    def test_qa_replays_returns_structured_citations(self):
        qa_record = DocumentQARecord(
            document_id=self.document.id,
            user_id=self.user.id,
            question="付款时间是什么？",
            answer="2026-07-01",
            source="document",
            citations='[{"page_number":3,"source_text":"应于2026年7月1日前支付。"}]',
            hit_chunks='[{"chunk_id":10,"content":"应于2026年7月1日前支付。"}]',
        )
        self.db.add(qa_record)
        self.db.commit()

        response = self.client.get(
            "/api/analytics/qa-replays",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["items"][0]["citations"][0]["page_number"], 3)
        self.assertEqual(payload["data"]["items"][0]["hit_chunks"][0]["chunk_id"], 10)

    def test_llm_billing_stats_returns_cost_summary(self):
        self.db.add(
            LLMCallLog(
                user_id=self.user.id,
                module_name="document",
                action="rag_answer",
                model_name="qwen-plus",
                input_tokens=1000,
                output_tokens=500,
                status="success",
            )
        )
        self.db.commit()
        response = self.client.get(
            "/api/analytics/llm-billing/stats",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertIn("summary", payload["data"])
        self.assertIn("pricing", payload["data"])

    def test_llm_routing_stats_requires_admin_and_returns_aggregate(self):
        self.db.add(
            LLMCallLog(
                user_id=self.user.id,
                module_name="chat",
                action="email_polish",
                model_name="qwen-turbo",
                status="success",
                request_id="routing-api-test",
                routing_role="small",
                routing_stage="initial",
            )
        )
        self.db.commit()

        forbidden = self.client.get("/api/analytics/llm-routing/stats", headers=self.headers)
        self.assertEqual(forbidden.status_code, 403)

        admin = User(
            username="admin_llm_routing",
            email="admin_llm_routing@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        response = self.client.get(
            "/api/analytics/llm-routing/stats",
            headers={"Authorization": f"Bearer {create_access_token({'sub': admin.id})}"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["data"]["routed_requests"], 1)

    def test_llm_routing_health_requires_admin(self):
        forbidden = self.client.get("/api/analytics/llm-routing/health", headers=self.headers)
        self.assertEqual(forbidden.status_code, 403)

        admin = User(
            username="admin_llm_routing_health",
            email="admin_llm_routing_health@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        response = self.client.get(
            "/api/analytics/llm-routing/health?hours=1",
            headers={"Authorization": f"Bearer {create_access_token({'sub': admin.id})}"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["data"]["status"], "ok")

    def test_database_module_supports_sqlite_connect_args(self):
        sqlite_kwargs = database_module.get_engine_kwargs("sqlite+pysqlite:///:memory:")
        mysql_kwargs = database_module.get_engine_kwargs("mysql+pymysql://root:123456@localhost:3306/aibg")

        self.assertEqual(sqlite_kwargs["connect_args"]["check_same_thread"], False)
        self.assertNotIn("pool_pre_ping", sqlite_kwargs)
        self.assertTrue(mysql_kwargs["pool_pre_ping"])

    def test_prompt_routes_require_admin_permission(self):
        response = self.client.get(
            "/api/prompts/",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 403)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "ADMIN_REQUIRED")
        self.assertEqual(payload["message"], "需要管理员权限")

    def test_prompt_create_returns_stable_schema_error_code(self):
        admin = User(
            username="admin_user",
            email="admin@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        self.db.refresh(admin)
        admin_token = create_access_token({"sub": admin.id})

        response = self.client.post(
            "/api/prompts/",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={
                "name": "bad_prompt",
                "template": "Question: {question}\nContext: {context}",
                "variables": "question",
                "change_note": "init bad prompt",
            },
        )

        self.assertEqual(response.status_code, 400)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "PROMPT_VARIABLE_SCHEMA_INVALID")

    def test_admin_can_start_prompt_rollout_and_rollback(self):
        admin = User(
            username="admin_release",
            email="admin_release@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        self.db.refresh(admin)
        admin_token = create_access_token({"sub": admin.id})

        tmpl = PromptTemplate(
            name="email_generate",
            description="邮件生成",
            variables="purpose",
        )
        self.db.add(tmpl)
        self.db.commit()
        self.db.refresh(tmpl)
        v1 = PromptTemplateVersion(
            template_id=tmpl.id,
            version=1,
            template="v1 -> {purpose}",
            is_active=True,
            change_note="init",
        )
        self.db.add(v1)
        self.db.commit()
        self.db.refresh(v1)
        tmpl.active_version_id = v1.id
        self.db.add(tmpl)
        self.db.commit()
        v2 = PromptTemplateVersion(
            template_id=tmpl.id,
            version=2,
            template="v2 -> {purpose}",
            is_active=False,
            change_note="candidate",
        )
        self.db.add(v2)
        self.db.commit()
        self.db.refresh(v2)

        rollout_response = self.client.post(
            f"/api/prompts/{tmpl.id}/rollout",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"version_id": v2.id, "rollout_percentage": 25},
        )

        self.assertEqual(rollout_response.status_code, 200)
        rollout_payload = rollout_response.json()
        self.assertTrue(rollout_payload["success"])
        self.assertEqual(rollout_payload["data"]["rollout"]["version_id"], v2.id)
        self.assertEqual(rollout_payload["data"]["rollout"]["percentage"], 25)

        rollback_response = self.client.post(
            f"/api/prompts/{tmpl.id}/rollback",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={},
        )

        self.assertEqual(rollback_response.status_code, 200)
        rollback_payload = rollback_response.json()
        self.assertTrue(rollback_payload["success"])
        self.assertIsNone(rollback_payload["data"]["rollout"])

    def test_prompt_rollback_returns_stable_error_when_unavailable(self):
        admin = User(
            username="admin_no_rollback",
            email="admin_no_rollback@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        self.db.refresh(admin)
        admin_token = create_access_token({"sub": admin.id})

        tmpl = PromptTemplate(
            name="meeting_summary",
            description="会议总结",
            variables="meeting_content",
        )
        self.db.add(tmpl)
        self.db.commit()
        self.db.refresh(tmpl)
        v1 = PromptTemplateVersion(
            template_id=tmpl.id,
            version=1,
            template="v1 -> {meeting_content}",
            is_active=True,
            change_note="init",
        )
        self.db.add(v1)
        self.db.commit()
        self.db.refresh(v1)
        tmpl.active_version_id = v1.id
        self.db.add(tmpl)
        self.db.commit()

        response = self.client.post(
            f"/api/prompts/{tmpl.id}/rollback",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={},
        )

        self.assertEqual(response.status_code, 400)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "PROMPT_ROLLBACK_NOT_AVAILABLE")

    def test_mcp_agent_types_returns_expected_contract(self):
        response = self.client.get(
            "/api/mcp/agent-types",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertGreaterEqual(payload["data"]["total"], 1)
        self.assertIn("agent_type", payload["data"]["items"][0])
        self.assertIn("allowed_tools", payload["data"]["items"][0])

    def test_agent_registry_returns_only_canonical_role_contracts(self):
        response = self.client.get("/api/agent/registry", headers=self.headers)

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["registry_version"], "enterprise_experts_v1")
        self.assertEqual(payload["data"]["task_protocol_version"], "agent_task_v1")
        self.assertEqual(
            [item["agent_type"] for item in payload["data"]["items"]],
            ["knowledge_agent", "legal_compliance_agent", "workflow_agent"],
        )
        self.assertEqual(payload["data"]["supervisor"]["agent_type"], "supervisor_agent")

    def test_mcp_tools_list_returns_filtered_specs(self):
        fake_tools = {
            "task_create_tool": FakeTool("task_create_tool", "创建任务"),
            "task_query_tool": FakeTool("task_query_tool", "查询任务"),
        }

        with patch.dict("app.mcp.registry._TOOL_INSTANCES", fake_tools, clear=True):
            response = self.client.get(
                "/api/mcp/tools",
                headers=self.headers,
                params={"agent_type": "workflow_agent"},
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["agent_type"], "workflow_agent")
        self.assertEqual([item["name"] for item in payload["data"]["items"]], ["task_create_tool", "task_query_tool"])

    def test_mcp_tool_call_returns_expected_contract(self):
        fake_tools = {
            "task_query_tool": FakeTool(
                "task_query_tool",
                "查询任务",
                auto_context_fields=("user_id",),
                parameters={
                    "type": "object",
                    "properties": {
                        "status": {"type": "string"},
                        "user_id": {"type": "integer"},
                    },
                    "required": ["status", "user_id"],
                },
                handler=lambda **kwargs: tool_success(
                    "查询成功",
                    {"status": kwargs["status"], "user_id": kwargs["user_id"]},
                ),
            )
        }

        with patch.dict("app.mcp.registry._TOOL_INSTANCES", fake_tools, clear=True):
            response = self.client.post(
                "/api/mcp/tools/call",
                headers=self.headers,
                json={
                    "tool_name": "task_query_tool",
                    "agent_type": "task_agent",
                    "arguments": {"status": "todo"},
                },
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["tool_name"], "task_query_tool")
        self.assertEqual(payload["data"]["agent_type"], "task_agent")
        self.assertTrue(payload["data"]["result"]["success"])
        self.assertEqual(payload["data"]["result"]["data"]["status"], "todo")
        self.assertEqual(payload["data"]["result"]["data"]["user_id"], self.user.id)

    def test_mcp_tool_call_hides_tool_error_detail(self):
        fake_tools = {
            "task_query_tool": FakeTool(
                "task_query_tool",
                "查询任务",
                auto_context_fields=("user_id",),
                parameters={
                    "type": "object",
                    "properties": {
                        "user_id": {"type": "integer"},
                    },
                    "required": ["user_id"],
                },
                handler=lambda **kwargs: {
                    "success": False,
                    "message": "任务查询失败",
                    "data": {"user_id": kwargs["user_id"]},
                    "error": "db_password=secret",
                },
            )
        }

        with patch.dict("app.mcp.registry._TOOL_INSTANCES", fake_tools, clear=True):
            response = self.client.post(
                "/api/mcp/tools/call",
                headers=self.headers,
                json={
                    "tool_name": "task_query_tool",
                    "agent_type": "task_agent",
                    "arguments": {},
                },
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertFalse(payload["data"]["result"]["success"])
        self.assertEqual(payload["data"]["result"]["message"], "任务查询失败")
        self.assertEqual(payload["data"]["result"]["error"], "任务查询失败")
        self.assertNotIn("secret", json.dumps(payload, ensure_ascii=False))

    def test_mcp_tool_call_rejects_sql_tool_for_non_admin(self):
        fake_tools = {
            "sql_query_tool": FakeTool("sql_query_tool", "执行 SQL"),
        }

        with patch.dict("app.mcp.registry._TOOL_INSTANCES", fake_tools, clear=True):
            response = self.client.post(
                "/api/mcp/tools/call",
                headers=self.headers,
                json={
                    "tool_name": "sql_query_tool",
                    "agent_type": "document_agent",
                    "arguments": {},
                },
            )

        self.assertEqual(response.status_code, 403)
        payload = response.json()
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["code"], "ADMIN_REQUIRED")

    def test_document_batch_upload_supports_governance_fields(self):
        with tempfile.TemporaryDirectory() as tmpdir, patch(
            "app.services.documents.document_service.UPLOAD_DIR",
            Path(tmpdir),
        ), patch(
            "app.services.documents.document_service._extract_segments",
            return_value=[{"text": "付款条款", "page_number": 1, "section_title": "正文"}],
        ), patch(
            "app.services.documents.document_service._split_text",
            return_value=[{"chunk_index": 0, "content": "付款条款", "page_number": 1, "section_title": "正文"}],
        ), patch(
            "app.services.documents.document_service._try_index_document",
            return_value=None,
        ):
            response = self.client.post(
                "/api/documents/batch-upload?async_mode=false",
                headers=self.headers,
                data={
                    "knowledge_base_name": "合同库",
                    "knowledge_base_category": "contract",
                    "classification": "legal",
                    "tags": "合同,付款",
                    "permission_scope": "restricted",
                    "permission_users": str(self.user.id),
                    "permission_roles": "admin",
                },
                files=[
                    ("files", ("a.md", b"alpha", "text/markdown")),
                    ("files", ("b.md", b"beta", "text/markdown")),
                ],
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["count"], 2)
        self.assertEqual(payload["data"]["documents"][0]["knowledge_base_id"], 1)
        self.assertEqual(payload["data"]["documents"][0]["classification"], "legal")
        self.assertEqual(payload["data"]["documents"][0]["permission_scope"], "restricted")

    def test_document_versions_returns_previous_versions(self):
        root = Document(
            user_id=self.user.id,
            title="合同.md",
            file_path="uploads/root.md",
            file_type="md",
            status="indexed",
            version_number=1,
            content_hash="hash-v1",
        )
        self.db.add(root)
        self.db.commit()
        self.db.refresh(root)
        v2 = Document(
            user_id=self.user.id,
            title="合同.md",
            file_path="uploads/v2.md",
            file_type="md",
            status="indexed",
            parent_document_id=root.id,
            version_number=2,
            content_hash="hash-v2",
        )
        self.db.add(v2)
        self.db.commit()

        response = self.client.get(f"/api/documents/{root.id}/versions", headers=self.headers)

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["total"], 2)
        self.assertEqual(payload["data"]["items"][0]["version_number"], 2)

    def test_mcp_tool_call_returns_approval_required_for_high_risk_tool(self):
        fake_tools = {
            "task_create_tool": FakeTool(
                "task_create_tool",
                "创建任务",
                auto_context_fields=("user_id",),
                parameters={
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "user_id": {"type": "integer"},
                    },
                    "required": ["title", "user_id"],
                },
            )
        }

        with patch.dict("app.mcp.registry._TOOL_INSTANCES", fake_tools, clear=True):
            response = self.client.post(
                "/api/mcp/tools/call",
                headers=self.headers,
                json={
                    "tool_name": "task_create_tool",
                    "agent_type": "task_agent",
                    "arguments": {"title": "审批测试任务"},
                },
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertFalse(payload["data"]["result"]["success"])
        self.assertEqual(payload["data"]["result"]["mcp_error_code"], "MCP_APPROVAL_REQUIRED")
        self.assertTrue(payload["data"]["result"]["data"]["approval_required"])

    def test_agent_approval_decision_endpoint_updates_status(self):
        from app.models.agent import AgentApprovalRequest

        approval = AgentApprovalRequest(
            user_id=self.user.id,
            tool_name="task_create_tool",
            agent_type="task_agent",
            approval_token="token-1",
            status="pending",
            risk_level="high",
        )
        self.db.add(approval)
        self.db.commit()
        self.db.refresh(approval)

        response = self.client.post(
            f"/api/agent/approvals/{approval.id}/decision",
            headers=self.headers,
            json={"approved": True, "decision_note": "允许执行"},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["status"], "approved")
        self.assertEqual(payload["data"]["decision_note"], "允许执行")

    def test_task_comments_and_logs_contract(self):
        task = Task(
            user_id=self.user.id,
            title="跟进合同",
            status="todo",
            priority="medium",
        )
        self.db.add(task)
        self.db.commit()
        self.db.refresh(task)

        comment_response = self.client.post(
            f"/api/tasks/{task.id}/comments",
            headers=self.headers,
            json={"content": "已联系法务同事协助审阅"},
        )
        update_response = self.client.put(
            f"/api/tasks/{task.id}",
            headers=self.headers,
            json={"progress": 40, "collaborators": ["法务", "采购"]},
        )
        comments_response = self.client.get(f"/api/tasks/{task.id}/comments", headers=self.headers)
        logs_response = self.client.get(f"/api/tasks/{task.id}/logs", headers=self.headers)

        self.assertEqual(comment_response.status_code, 200)
        self.assertEqual(update_response.status_code, 200)
        self.assertEqual(comments_response.status_code, 200)
        self.assertEqual(logs_response.status_code, 200)
        self.assertEqual(update_response.json()["data"]["progress"], 40)
        self.assertEqual(update_response.json()["data"]["collaborators"], ["法务", "采购"])
        self.assertEqual(comments_response.json()["data"][0]["content"], "已联系法务同事协助审阅")
        self.assertTrue(any(item["action"] == "comment_added" for item in logs_response.json()["data"]))

    def test_tool_health_and_feedback_bundle_contract(self):
        admin = User(
            username="admin_tool_health",
            email="admin_tool_health@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        self.db.refresh(admin)
        admin_token = create_access_token({"sub": admin.id})

        qa_record = DocumentQARecord(
            document_id=self.document.id,
            user_id=self.user.id,
            question="付款条件是什么？",
            answer="不知道",
            source="document",
            feedback_value="negative",
            feedback_status="open",
            feedback_reason="incorrect_answer",
            feedback_created_at=datetime.utcnow(),
        )
        self.db.add(qa_record)
        agent_run = AgentRun(user_id=self.user.id, goal="创建任务", status="completed")
        self.db.add(agent_run)
        self.db.commit()
        self.db.refresh(agent_run)
        self.db.add(
            ToolCallLog(
                agent_run_id=agent_run.id,
                step=1,
                tool_name="task_create_tool",
                status="pending_approval",
                error="工具调用需要人工审批",
                duration_ms=10,
            )
        )
        self.db.commit()

        tool_health_response = self.client.get("/api/analytics/tool-health", headers=self.headers)
        bundle_response = self.client.post(
            "/api/analytics/feedback/export-eval-bundle",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"days": 30},
        )

        self.assertEqual(tool_health_response.status_code, 200)
        self.assertEqual(bundle_response.status_code, 200)
        self.assertTrue(tool_health_response.json()["success"])
        self.assertEqual(tool_health_response.json()["data"]["items"][0]["tool_name"], "task_create_tool")
        self.assertTrue(bundle_response.json()["success"])
        self.assertGreaterEqual(bundle_response.json()["data"]["count"], 1)

    def test_org_and_department_management_contract(self):
        admin = User(
            username="admin_org",
            email="admin_org@example.com",
            hashed_password=hash_password("secret"),
            role="admin",
        )
        self.db.add(admin)
        self.db.commit()
        self.db.refresh(admin)
        admin_token = create_access_token({"sub": admin.id})

        org_response = self.client.post(
            "/api/org/organizations",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"name": "OpenAI CN", "code": "OA-CN", "description": "China org"},
        )
        self.assertEqual(org_response.status_code, 200)
        org_id = org_response.json()["data"]["id"]

        dept_response = self.client.post(
            "/api/org/departments",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"organization_id": org_id, "name": "Legal", "code": "LEGAL"},
        )
        self.assertEqual(dept_response.status_code, 200)
        dept_id = dept_response.json()["data"]["id"]

        assign_response = self.client.post(
            f"/api/org/users/{self.user.id}/assign",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"organization_id": org_id, "department_id": dept_id, "job_title": "PM"},
        )
        self.assertEqual(assign_response.status_code, 200)
        self.assertEqual(assign_response.json()["data"]["organization_id"], org_id)
        self.assertEqual(assign_response.json()["data"]["department_id"], dept_id)

        users_response = self.client.get(
            "/api/auth/users",
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        self.assertEqual(users_response.status_code, 200)
        self.assertTrue(users_response.json()["success"])
        self.assertTrue(any(item["id"] == self.user.id for item in users_response.json()["data"]))

    def test_task_list_supports_source_filters(self):
        document_task = Task(
            user_id=self.user.id,
            title="文档任务",
            status="todo",
            priority="medium",
            source_type="document",
            source_id=11,
        )
        meeting_task = Task(
            user_id=self.user.id,
            title="会议任务",
            status="todo",
            priority="medium",
            source_type="meeting",
            source_id=22,
        )
        self.db.add_all([document_task, meeting_task])
        self.db.commit()

        response = self.client.get(
            "/api/tasks/",
            headers=self.headers,
            params={"source_type": "document", "source_id": 11, "page": 1, "page_size": 20},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        ids = [item["id"] for item in payload["data"]["items"]]
        self.assertIn(document_task.id, ids)
        self.assertNotIn(meeting_task.id, ids)

    def test_agent_run_list_supports_artifact_filters(self):
        document_run = AgentRun(
            user_id=self.user.id,
            goal="分析文档风险",
            status="completed",
            result=json.dumps({"artifacts": {"documents": [{"document_id": 101, "summary": "合同风险"}]}}, ensure_ascii=False),
            final_answer="已完成文档分析",
            total_steps=2,
        )
        task_run = AgentRun(
            user_id=self.user.id,
            goal="创建任务",
            status="completed",
            result=json.dumps({"artifacts": {"tasks": [{"task_id": 202, "title": "跟进合同"}]}}, ensure_ascii=False),
            final_answer="已创建任务",
            total_steps=1,
        )
        self.db.add_all([document_run, task_run])
        self.db.commit()

        response = self.client.get(
            "/api/agent/runs",
            headers=self.headers,
            params={"artifact_type": "document", "artifact_id": 101, "page": 1, "page_size": 20},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        ids = [item["id"] for item in payload["data"]["items"]]
        self.assertIn(document_run.id, ids)
        self.assertNotIn(task_run.id, ids)

    def test_sensitive_document_filter_contract(self):
        sensitive_doc = Document(
            user_id=self.user.id,
            title="confidential-plan.md",
            file_path="uploads/confidential-plan.md",
            file_type="md",
            status="indexed",
            sensitivity_level="confidential",
        )
        self.db.add(sensitive_doc)
        self.db.commit()

        response = self.client.get(
            "/api/documents/",
            headers=self.headers,
            params={"sensitivity_level": "confidential"},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"]["items"][0]["sensitivity_level"], "confidential")

if __name__ == "__main__":
    unittest.main()
